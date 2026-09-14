import csv
import math
import re
from io import BytesIO
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .services import DomainError, role


def clean_text(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))[:32767]


def csv_safe(value):
    if not isinstance(value, str): return value
    value = clean_text(value)
    return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value


def table_file(columns, rows, file_format, filename='srez-table', title='Срез'):
    if file_format not in ['csv', 'xlsx']:
        raise DomainError('Выберите CSV или Excel (.xlsx).')
    if file_format == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response.write('\ufeff')
        writer = csv.writer(response, delimiter=';')
        writer.writerow([csv_safe(v) for v in columns])
        writer.writerows([[csv_safe(v) for v in row] for row in rows])
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = 'Данные'
        book.properties.title = clean_text(title)
        for row_number, values in enumerate([columns, *rows], 1):
            for column_number, value in enumerate(values, 1):
                cell = sheet.cell(row_number, column_number)
                if isinstance(value, str):
                    cell.value = clean_text(value)
                    # Source formulas and text beginning with = remain literal text.
                    cell.data_type = 's'
                else:
                    cell.value = value
                    if isinstance(value, float): cell.number_format = '#,##0.00'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if row_number == 1:
                    cell.font = Font(name='Arial', bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='244FAD')
                else:
                    cell.font = Font(name='Arial', size=11)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for i, name in enumerate(columns, 1):
            width = max([len(str(name)), *(len(str(row[i-1] or '')) for row in rows[:150])])
            sheet.column_dimensions[get_column_letter(i)].width = min(48, max(14, width + 2))
        output = BytesIO(); book.save(output)
        response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}.{file_format}"'
    return response


def export_table_data(request, data):
    if role(request.user) not in ['office', 'manager']:
        raise DomainError('Выгрузка доступна офису и руководителю.', 403)
    columns, rows = data.get('columns'), data.get('rows')
    if not isinstance(columns, list) or not 1 <= len(columns) <= 256 or any(not isinstance(c, str) or len(c) > 500 for c in columns):
        raise DomainError('Некорректные заголовки таблицы.')
    if not isinstance(rows, list) or len(rows) > 20000 or len(rows) * len(columns) > 250000:
        raise DomainError('Слишком большая выгрузка. Уточните фильтры.')
    for row in rows:
        if not isinstance(row, list) or len(row) != len(columns): raise DomainError('Некорректная строка таблицы.')
        for value in row:
            if value is not None and (type(value) not in (str, int, float) or isinstance(value, float) and not math.isfinite(value)):
                raise DomainError('Некорректное значение ячейки.')
            if isinstance(value, str) and len(value) > 32767: raise DomainError('Слишком длинное значение ячейки.')
    return table_file(columns, rows, data.get('format'), title=str(data.get('title', 'Срез'))[:200])
