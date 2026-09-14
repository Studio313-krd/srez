import csv
from io import BytesIO, StringIO
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from .models import Access


@override_settings(SECURE_SSL_REDIRECT=False)
class TableExportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser('export-admin')
        self.client.force_login(self.user)
        self.data = {'columns':['Магазин', 'Выручка', 'Чеки', 'Комментарий'],
            'rows':[['Магазин 10', 9978.91, 0, '=1+1'], ['Магазин 2', None, 10, '+cmd']],
            'format':'xlsx', 'title':'Отчёты', 'context':'2026-09-13-close'}

    def post(self):
        return self.client.post('/api/table-export/', self.data, content_type='application/json')

    def test_excel_preserves_order_numbers_zero_blank_and_literal_formulas(self):
        response = self.post()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response['Content-Disposition'].endswith('.xlsx"'))
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(list(sheet.values), [('Магазин','Выручка','Чеки','Комментарий'),
            ('Магазин 10', 9978.91, 0, '=1+1'), ('Магазин 2', None, 10, '+cmd')])
        self.assertEqual(sheet['D2'].data_type, 's')
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:D3')

    def test_csv_is_utf8_and_escapes_formula_injection(self):
        self.data['format'] = 'csv'
        response = self.post()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'\xef\xbb\xbf'))
        rows = list(csv.reader(StringIO(response.content.decode('utf-8-sig')), delimiter=';'))
        self.assertEqual(rows[1], ['Магазин 10','9978.91','0',"'=1+1"])
        self.assertEqual(rows[2][-1], "'+cmd")

    def test_missing_or_invalid_format_is_rejected(self):
        for fmt in [None, '', 'xls', 'pdf']:
            self.data['format'] = fmt
            self.assertEqual(self.post().status_code, 400)

    def test_invalid_cells_and_ragged_rows_are_rejected(self):
        for rows in [[['short']], [[{},0,0,'x']], [[float('nan'),0,0,'x']]]:
            self.data['rows'] = rows
            self.assertEqual(self.post().status_code, 400)

    def test_seller_and_anonymous_cannot_export_office_tables(self):
        seller = get_user_model().objects.create_user('export-seller')
        Access.objects.create(user=seller, role='store')
        self.client.force_login(seller)
        self.assertEqual(self.post().status_code, 403)
        self.client.logout()
        self.assertEqual(self.post().status_code, 401)

    def test_empty_result_still_has_headers(self):
        self.data['rows'] = []
        sheet = load_workbook(BytesIO(self.post().content)).active
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual(sheet.max_column, 4)
