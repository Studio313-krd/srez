import json
from pathlib import Path
from shutil import copy2
import openpyxl

source = Path(r'C:\Users\dev\Downloads\Telegram Desktop\Ред02_09_18_26_Кадровый_реестр_торговой_сети_Мильстрим_и_Культура.xlsx')
destination = Path(__file__).resolve().parent.parent / '.local/import/sources'
destination.mkdir(parents=True, exist_ok=True)
copy2(source, destination / 'registry.xlsx')
workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
formulas = openpyxl.load_workbook(source, data_only=False, read_only=True)
result = {'path': str(source), 'title': source.name, 'sheets': []}
for sheet in workbook:
    rows = list(sheet.iter_rows(values_only=True))
    result['sheets'].append({'title': sheet.title, 'values': rows,
        'formulas': [[c.value if c.data_type == 'f' else None for c in row] for row in formulas[sheet.title]]})
    print(json.dumps({'sheet': sheet.title, 'rows': len(rows), 'header': rows[2:5]}, ensure_ascii=False, default=str))
(destination / 'hr.json').write_text(json.dumps(result, ensure_ascii=False, default=str), encoding='utf-8')
